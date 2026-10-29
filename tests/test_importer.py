import io

import pandas as pd

from app.importer import identify_columns, import_poles_file, ler_coordenadas, limpar_codigo
from app.servicos.dados_postes import (
    atualizar_status_poste,
    definir_localizacao_poste,
    obter_poste_por_codigo,
)


def _csv(linhas):
    buffer = io.StringIO()
    pd.DataFrame(linhas).to_csv(buffer, index=False)
    return buffer.getvalue().encode("utf-8")


# --- detecção de colunas -----------------------------------------------------

def test_tipo_poste_nao_vira_codigo():
    mapa = identify_columns(["Plaqueta", "Tipo Poste", "Rua"])
    assert mapa["codigo"] == "Plaqueta"
    assert mapa["tipo_poste"] == "Tipo Poste"


def test_ponto_de_referencia_nao_vira_codigo():
    mapa = identify_columns(["Código", "Ponto de Referência"])
    assert mapa["codigo"] == "Código"
    assert mapa["referencia"] == "Ponto de Referência"


def test_sinonimo_curto_so_casa_exato():
    # "x", "y" e "w" não podem casar com qualquer coluna que contenha a letra
    mapa = identify_columns(["codigo", "taxa_y", "complexo", "wifi"])
    assert "latitude" not in mapa
    assert "longitude" not in mapa
    assert "potencia" not in mapa


def test_exportacao_kobo():
    colunas = [
        "start", "end", "Código da plaqueta", "Rua", "Bairro", "Tipo de luminária",
        "Potência", "localizacao", "_localizacao_latitude", "_localizacao_longitude",
        "_localizacao_altitude", "_localizacao_precision", "Foto do poste", "_id", "_uuid",
    ]
    mapa = identify_columns(colunas)
    assert mapa["codigo"] == "Código da plaqueta"
    assert mapa["latitude"] == "_localizacao_latitude"
    assert mapa["longitude"] == "_localizacao_longitude"
    assert mapa["logradouro"] == "Rua"
    assert mapa["tipo_luminaria"] == "Tipo de luminária"
    assert mapa["potencia"] == "Potência"


def test_coluna_nao_e_usada_duas_vezes():
    mapa = identify_columns(["codigo", "poste"])
    assert list(mapa.values()).count("codigo") <= 1
    assert mapa["codigo"] == "codigo"


# --- limpeza de valores ------------------------------------------------------

def test_codigo_numerico_do_excel():
    assert limpar_codigo(1001.0) == "1001"
    assert limpar_codigo(" P-01 ") == "P-01"
    assert limpar_codigo(float("nan")) is None


def test_coordenadas():
    assert ler_coordenadas("-6,15", "-35,60")[:2] == (-6.15, -35.6)
    assert ler_coordenadas(None, None) == (None, None, None)
    assert ler_coordenadas("abc", "-35")[0] is None
    assert ler_coordenadas(200, -35)[0] is None
    assert ler_coordenadas(0, 0)[0] is None


# --- gravação ----------------------------------------------------------------

def test_reimportar_preserva_status_e_localizacao():
    import_poles_file(_csv([{"codigo": "P-01", "rua": "Rua A", "bairro": "Centro"}]), "a.csv")
    atualizar_status_poste("P-01", "urgente")
    definir_localizacao_poste("P-01", -6.15, -35.61)

    resultado = import_poles_file(_csv([{"codigo": "P-01", "rua": "Rua Nova", "bairro": "Centro"}]), "b.csv")

    poste = obter_poste_por_codigo("P-01")
    assert resultado["importados"] == 1
    assert poste["logradouro"] == "Rua Nova"
    assert poste["status"] == "urgente"
    assert (poste["latitude"], poste["longitude"]) == (-6.15, -35.61)


def test_reimportar_com_coordenada_valida_atualiza():
    import_poles_file(_csv([{"codigo": "P-02", "latitude": -6.1, "longitude": -35.5}]), "a.csv")
    import_poles_file(_csv([{"codigo": "P-02", "latitude": -6.2, "longitude": -35.7}]), "b.csv")
    poste = obter_poste_por_codigo("P-02")
    assert (poste["latitude"], poste["longitude"]) == (-6.2, -35.7)


def test_celula_vazia_nao_apaga_dado_existente():
    import_poles_file(_csv([{"codigo": "P-03", "rua": "Rua A", "bairro": "Bela Vista"}]), "a.csv")
    import_poles_file(_csv([{"codigo": "P-03", "rua": "Rua B", "bairro": None}]), "b.csv")
    poste = obter_poste_por_codigo("P-03")
    assert poste["logradouro"] == "Rua B"
    assert poste["bairro"] == "Bela Vista"


def test_poste_novo_sem_coordenada_fica_pendente():
    import_poles_file(_csv([{"codigo": "P-04"}]), "a.csv")
    poste = obter_poste_por_codigo("P-04")
    assert poste["latitude"] is None
    assert poste["status"] == "normal"


def test_codigo_repetido_na_planilha():
    resultado = import_poles_file(_csv([{"codigo": "P-05", "rua": "A"}, {"codigo": "P-05", "rua": "B"}]), "a.csv")
    assert resultado["importados"] == 1
    assert resultado["duplicados"] == 1
    assert obter_poste_por_codigo("P-05")["logradouro"] == "B"


# --- exportação do KoboToolbox (formulário kobo/levantamento_postes.xlsx) -----

_KOBO_META = {"_id": 1, "_uuid": "abc", "_submission_time": "2026-10-27T10:00:00", "_validation_status": "",
              "_notes": "", "_status": "submitted_via_web", "_submitted_by": "", "__version__": "v1",
              "_tags": "", "_index": 1}


def _kobo_csv(linhas, sep=";"):
    buffer = io.StringIO()
    pd.DataFrame(linhas).to_csv(buffer, index=False, sep=sep)
    return buffer.getvalue().encode("utf-8")


def test_kobo_cabecalho_xml():
    linha = {
        "start": "2026-10-27T09:58:00", "end": "2026-10-27T10:00:00",
        "codigo": "123456", "localizacao": "-6.15 -35.61 120 4.5",
        "_localizacao_latitude": -6.15, "_localizacao_longitude": -35.61,
        "_localizacao_altitude": 120, "_localizacao_precision": 4.5,
        "foto": "foto.jpg", "foto_URL": "https://kf.kobotoolbox.org/x.jpg",
        "logradouro": "Rua A", "bairro": "Centro", "referencia": "UBS",
        "tipo_luminaria": "vapor_de_sodio", "potencia": 70, "tipo_poste": "concreto_circular",
        "tipo_braco": "braco_curto_1m", "observacoes": "Acesa de dia", **_KOBO_META,
    }
    resultado = import_poles_file(_kobo_csv([linha]), "kobo.csv")
    poste = obter_poste_por_codigo("123456")
    assert resultado["importados"] == 1
    assert resultado["colunas_identificadas"]["logradouro"] == "logradouro"  # e não a coluna "end"
    assert (poste["latitude"], poste["longitude"]) == (-6.15, -35.61)
    assert poste["tipo_luminaria"] == "Vapor de Sódio"
    assert poste["tipo_poste"] == "Concreto Circular"
    assert poste["tipo_braco"] == "Braço Curto 1m"
    assert poste["potencia"] == 70
    assert poste["observacoes"] == "Acesa de dia"


def test_kobo_cabecalho_rotulos():
    linha = {
        "start": "2026-10-27T09:58:00", "end": "2026-10-27T10:00:00",
        "Código da plaqueta": "654321", "Localização do poste": "-6.16 -35.62 0 3",
        "_Localização do poste_latitude": -6.16, "_Localização do poste_longitude": -35.62,
        "_Localização do poste_altitude": 0, "_Localização do poste_precision": 3,
        "Foto do poste": "foto.jpg", "Foto do poste_URL": "https://kf.kobotoolbox.org/y.jpg",
        "Rua / Logradouro": "Rua B", "Bairro ou comunidade": "Lagoa do Mato",
        "Ponto de referência": "Igreja", "Tipo de luminária": "LED", "Potência (W)": 100,
        "Tipo de estrutura": "Madeira Tratada", "Tipo de braço": "Topo de Poste",
        "Observações": "", **_KOBO_META,
    }
    resultado = import_poles_file(_kobo_csv([linha]), "kobo.csv")
    mapa = resultado["colunas_identificadas"]
    poste = obter_poste_por_codigo("654321")
    assert mapa["codigo"] == "Código da plaqueta"
    assert mapa["logradouro"] == "Rua / Logradouro"
    assert mapa["tipo_poste"] == "Tipo de estrutura"
    assert (poste["latitude"], poste["longitude"]) == (-6.16, -35.62)
    assert poste["bairro"] == "Lagoa do Mato"
    assert poste["tipo_braco"] == "Topo de Poste"


def test_formulario_kobo_usa_as_opcoes_do_sistema(tmp_path):
    from openpyxl import load_workbook

    from app.constantes import BRACOS, ESTRUTURAS, LUMINARIAS
    from kobo.gerar_formulario import gerar

    wb = load_workbook(gerar(tmp_path / "form.xlsx"))
    rotulos = {linha[2] for linha in wb["choices"].iter_rows(min_row=2, values_only=True)}
    assert set(LUMINARIAS + ESTRUTURAS + BRACOS) == rotulos
    nomes = {linha[1] for linha in wb["survey"].iter_rows(min_row=2, values_only=True)}
    assert {"codigo", "localizacao", "logradouro", "bairro", "tipo_luminaria", "potencia"} <= nomes


# --- correções da auditoria (out/2026) -----------------------------------------

def _xlsx(linhas):
    buffer = io.BytesIO()
    pd.DataFrame(linhas).to_excel(buffer, index=False)
    return buffer.getvalue()


def test_xlsx_preserva_zero_a_esquerda():
    import_poles_file(_xlsx([{"codigo": "00123", "rua": "Rua A"}]), "a.xlsx")
    assert obter_poste_por_codigo("00123") is not None


def test_xls_antigo_recusado_com_mensagem_clara():
    import pytest

    with pytest.raises(ValueError, match="salve como .xlsx"):
        import_poles_file(b"\xD0\xCF\x11\xE0" + b"\x00" * 60, "kobo.xls")


def test_coluna_geopoint_unica_do_kobo():
    for cabecalho in ("localizacao", "Localização do poste"):
        res = import_poles_file(
            _csv([{"Código da plaqueta": f"G-{cabecalho[:3]}", cabecalho: "-6.1650 -35.6010 52.0 4.5"}]), "k.csv"
        )
        poste = obter_poste_por_codigo(f"G-{cabecalho[:3]}")
        assert (poste["latitude"], poste["longitude"]) == (-6.165, -35.601), res
        assert res["colunas_identificadas"]["codigo"] == "Código da plaqueta"


def test_rotulo_com_poste_nao_vira_codigo():
    mapa = identify_columns(["Código da plaqueta", "Localização do poste", "Foto do poste"])
    assert mapa["codigo"] == "Código da plaqueta"
    assert identify_columns(["Foto do poste", "Rua"]).get("codigo") is None


def test_avisa_quando_planilha_nao_tem_coordenada():
    res = import_poles_file(_csv([{"codigo": "S-1", "rua": "Rua A"}]), "a.csv")
    assert any("localização pendente" in a for a in res["detalhes_erros"])


def test_coordenada_trocada_e_corrigida_e_distante_ignorada():
    assert ler_coordenadas("-35.60", "-6.16")[:2] == (-6.16, -35.60)
    lat, lng, aviso = ler_coordenadas("-23.55", "-46.63")  # São Paulo
    assert (lat, lng) == (None, None) and "km" in aviso


def test_potencia_milhar_e_invalida():
    from app.importer import ler_potencia

    assert ler_potencia("1.000") == 1000
    assert ler_potencia("100 W") == 100
    assert ler_potencia("70,5") == 70
    assert ler_potencia("-50") is None and ler_potencia("0") is None


def test_busca_trata_curinga_como_texto():
    from app.servicos.dados_postes import cadastrar_poste, listar_postes

    cadastrar_poste({"codigo": "P_01", "logradouro": "Rua A"})
    cadastrar_poste({"codigo": "PX01", "logradouro": "Rua B"})
    assert [p["codigo"] for p in listar_postes(termo_busca="P_01")] == ["P_01"]
    assert listar_postes(termo_busca="100%") == []
