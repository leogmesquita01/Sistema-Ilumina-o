"""Área do administrador no modo local (SQLite) e visibilidade da tela."""
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
import io

from app.constantes import OS_ABERTA, OS_CANCELADA, STATUS_NORMAL, STATUS_URGENTE
from app.servicos import admin_service as adm
from app.servicos.dados_postes import PosteDuplicado, cadastrar_poste, obter_poste_por_codigo
from app.servicos.ordens_service import criar_ordem_servico, listar_ordens

APP = str(Path(__file__).resolve().parent.parent / "app" / "dashboard" / "main.py")


@pytest.fixture
def base():
    cadastrar_poste({"codigo": "P-01"})
    cadastrar_poste({"codigo": "P-02"})
    return criar_ordem_servico("P-01", "Lâmpada apagada", prioridade="Urgente")


def test_renomear_poste_leva_as_os_junto(base):
    adm.renomear_poste("p-01", "P-001")
    assert obter_poste_por_codigo("P-01") is None
    assert [o["codigo_poste"] for o in listar_ordens("Todos")] == ["P-001"]


def test_renomear_para_codigo_existente_recusa(base):
    with pytest.raises(PosteDuplicado):
        adm.renomear_poste("P-01", "p-02")


def test_renomear_so_a_caixa_e_permitido(base):
    adm.renomear_poste("P-02", "p-02")
    assert obter_poste_por_codigo("p-02")["codigo"] == "p-02"


def test_excluir_poste_com_os_recusa_e_sem_os_exclui(base):
    with pytest.raises(adm.AcaoRecusada):
        adm.excluir_poste("P-01")
    adm.excluir_poste("P-02")
    assert obter_poste_por_codigo("P-02") is None


def test_cancelar_os_volta_poste_ao_normal(base):
    assert obter_poste_por_codigo("P-01")["status"] == STATUS_URGENTE
    adm.cancelar_ordem(base, "aberta em duplicidade")
    ordem = listar_ordens("Todos")[0]
    assert ordem["status"] == OS_CANCELADA
    assert "duplicidade" in ordem["observacoes"]
    assert obter_poste_por_codigo("P-01")["status"] == STATUS_NORMAL


def test_cancelar_exige_motivo_e_os_aberta(base):
    with pytest.raises(adm.AcaoRecusada):
        adm.cancelar_ordem(base, "  ")
    adm.cancelar_ordem(base, "teste")
    with pytest.raises(adm.AcaoRecusada):
        adm.cancelar_ordem(base, "de novo")


def test_corrigir_protocolo(base):
    adm.corrigir_protocolo(base, "OS-2026-9999")
    assert listar_ordens("Todos")[0]["protocolo"] == "OS-2026-9999"
    outro = criar_ordem_servico("P-02", "Braço quebrado")
    with pytest.raises(adm.AcaoRecusada):
        adm.corrigir_protocolo(outro, "OS-2026-9999")


def test_excluir_os_recalcula_poste(base):
    adm.excluir_ordem(base)
    assert listar_ordens("Todos") == []
    assert obter_poste_por_codigo("P-01")["status"] == STATUS_NORMAL


def test_backup_tem_postes_e_ordens(base):
    livro = load_workbook(io.BytesIO(adm.gerar_backup_xlsx()))
    assert livro.sheetnames == ["Postes", "Ordens de serviço"]
    assert livro["Postes"].max_row == 3  # cabeçalho + 2 postes
    assert livro["Ordens de serviço"].max_row == 2


def test_resumir_mudanca():
    texto = adm.resumir_mudanca({"status": OS_ABERTA, "data_atualizacao": 1}, {"status": OS_CANCELADA, "data_atualizacao": 2})
    assert texto == "status: aberta → cancelada"


def test_equipe_e_historico_vazios_no_modo_local():
    assert adm.listar_equipe() == []
    assert adm.listar_historico() == []


# ---- Tela --------------------------------------------------------------------

def _menu(at):
    return at.radio(key="nav_principal").options


def test_admin_ve_a_tela(base):
    at = AppTest.from_file(APP, default_timeout=120)
    at.session_state["pagina_atual"] = "admin"
    at.run()
    assert not at.exception, at.exception
    assert any("Admin" in o for o in _menu(at))
    assert any("Administração" in m.value for m in at.markdown)


def test_operador_nao_ve_a_tela(base, monkeypatch):
    from app.dashboard import sessao

    monkeypatch.setattr(sessao, "_USUARIO_LOCAL", {**sessao._USUARIO_LOCAL, "papel": "operador"})
    at = AppTest.from_file(APP, default_timeout=120)
    at.session_state["pagina_atual"] = "admin"  # tentativa de entrar direto
    at.run()
    assert not at.exception, at.exception
    assert not any("Admin" in o for o in _menu(at))
    assert at.session_state["pagina_atual"] == "mapa"


def test_cancelar_pela_tela(base):
    at = AppTest.from_file(APP, default_timeout=120)
    at.session_state["pagina_atual"] = "admin"
    at.run()
    motivo = next(t for t in at.text_input if t.label == "Motivo")
    motivo.input("lançada no poste errado")
    next(b for b in at.button if b.label == "Cancelar O.S.").click()
    at.run()
    assert not at.exception, at.exception
    assert any("cancelada" in s.value for s in at.success)
    assert listar_ordens("Todos")[0]["status"] == OS_CANCELADA
