"""Gera o formulário do KoboToolbox (XLSForm) para o levantamento de postes.

Uso:  python kobo/gerar_formulario.py
Saída: kobo/levantamento_postes.xlsx — envie no Kobo em "Novo projeto → Enviar XLSForm".

Os nomes dos campos e as opções saem das mesmas constantes do sistema, então a
exportação do Kobo (Excel/CSV) entra direto na tela "Importar planilha".
"""
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from app.constantes import BRACOS, ESTRUTURAS, LUMINARIAS  # noqa: E402
from app.importer import normalizar_nome_coluna  # noqa: E402

SAIDA = Path(__file__).resolve().parent / "levantamento_postes.xlsx"

CABECALHO = ["type", "name", "label", "hint", "required", "appearance", "parameters",
             "default", "constraint", "constraint_message"]

SURVEY = [
    # Metadados que o Kobo preenche sozinho (hora de início/fim da coleta).
    {"type": "start", "name": "start"},
    {"type": "end", "name": "end"},
    {"type": "text", "name": "codigo", "label": "Código da plaqueta", "required": "yes",
     "hint": "Exatamente como está na plaqueta da COSERN"},
    {"type": "geopoint", "name": "localizacao", "label": "Localização do poste", "required": "yes",
     "hint": "Fique ao lado do poste e espere a precisão ficar abaixo de 10 m. Funciona sem internet.",
     "parameters": "capture-accuracy=10 warning-accuracy=20"},
    {"type": "image", "name": "foto", "label": "Foto do poste",
     "hint": "Se possível, mostrando a plaqueta e a luminária",
     # Foto reduzida: ~200 KB em vez de 3-5 MB (cabe no plano gratuito do Kobo).
     "parameters": "max-pixels=1024"},
    {"type": "text", "name": "logradouro", "label": "Rua / Logradouro"},
    {"type": "text", "name": "bairro", "label": "Bairro ou comunidade", "required": "yes"},
    {"type": "text", "name": "referencia", "label": "Ponto de referência",
     "hint": "Ex.: em frente à UBS, ao lado da igreja"},
    {"type": "select_one luminaria", "name": "tipo_luminaria", "label": "Tipo de luminária", "required": "yes"},
    {"type": "integer", "name": "potencia", "label": "Potência (W)", "default": "100",
     "hint": "Se não souber, deixe 100",
     "constraint": ". > 0 and . <= 1000", "constraint_message": "Valor entre 1 e 1000 W"},
    {"type": "select_one estrutura", "name": "tipo_poste", "label": "Tipo de estrutura"},
    {"type": "select_one braco", "name": "tipo_braco", "label": "Tipo de braço"},
    {"type": "text", "name": "observacoes", "label": "Observações", "appearance": "multiline",
     "hint": "Lâmpada apagada, acesa de dia, fiação solta, etc."},
]

LISTAS = {"luminaria": LUMINARIAS, "estrutura": ESTRUTURAS, "braco": BRACOS}

SETTINGS = [
    ["form_title", "form_id", "version", "style"],
    ["Ilumina - Levantamento de Postes", "ilumina_postes", "2026102601", "pages"],
]


def gerar(destino: Path = SAIDA) -> Path:
    wb = Workbook()

    survey = wb.active
    survey.title = "survey"
    survey.append(CABECALHO)
    for campo in SURVEY:
        survey.append([campo.get(col, "") for col in CABECALHO])

    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label"])
    for lista, opcoes in LISTAS.items():
        for opcao in opcoes:
            # name sem acento/espaço; o importador converte de volta para a opção oficial.
            choices.append([lista, normalizar_nome_coluna(opcao), opcao])

    settings = wb.create_sheet("settings")
    for linha in SETTINGS:
        settings.append(linha)

    for ws in wb.worksheets:
        for celula in ws[1]:
            celula.font = Font(bold=True)
        for coluna in ws.columns:
            largura = max(len(str(c.value or "")) for c in coluna) + 2
            ws.column_dimensions[coluna[0].column_letter].width = min(largura, 60)

    wb.save(destino)
    return destino


if __name__ == "__main__":
    print(f"Formulário gerado em {gerar()}")
